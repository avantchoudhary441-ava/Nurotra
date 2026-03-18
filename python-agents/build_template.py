import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

def create_styled_template():
    # Load default template
    prs = Presentation()
    
    for layout in prs.slide_layouts:
        background = layout.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(15, 15, 15) # Very dark gray/black
        
        # Style placeholders
        for shape in layout.placeholders:
            if shape.has_text_frame:
                if shape.name.lower().startswith("title") or "heading" in shape.name.lower():
                    # Cyan accent for titles
                    for paragraph in shape.text_frame.paragraphs:
                        try:
                            paragraph.font.color.rgb = RGBColor(0, 255, 255) # Cyan
                            paragraph.font.name = "Arial"
                        except:
                            pass
                else:
                    # Light Gray for body text
                    for paragraph in shape.text_frame.paragraphs:
                        try:
                            paragraph.font.color.rgb = RGBColor(220, 220, 220)
                            paragraph.font.name = "Arial"
                        except:
                            pass
                        
    # Save as Nurotra Template
    template_path = os.path.join(os.path.dirname(__file__), 'nurotra_template.pptx')
    prs.save(template_path)
    print(f"Successfully created template at {template_path}")

if __name__ == '__main__':
    create_styled_template()
