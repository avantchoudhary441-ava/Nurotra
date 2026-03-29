import os
import json
import base64
import tempfile
import sys
import requests
import random
from io import BytesIO
from typing import Dict, Any, List, Optional, Tuple

# CAMEL dependencies
from camel.agents import ChatAgent
from camel.messages import BaseMessage
from camel.models import ModelFactory
from camel.types import ModelPlatformType
from camel.configs import ChatGPTConfig

# PPTX rendering
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE, MSO_AUTO_SHAPE_TYPE

# ==========================================
# 1. API KEYS & TOOLS
# ==========================================
PEXELS_API_KEY = os.getenv("PEXELS_API_KEY")
UNSPLASH_ACCESS_KEY = os.getenv("UNSPLASH_ACCESS_KEY")

def search_image(query: str) -> Optional[BytesIO]:
    """Search for an image using Pexels first, then Unsplash fallback."""
    if not query or query.lower() in ["none", "null"]: return None
    # 1. Try Pexels
    if PEXELS_API_KEY:
        try:
            url = f"https://api.pexels.com/v1/search?query={query}&per_page=1"
            headers = {"Authorization": PEXELS_API_KEY}
            resp = requests.get(url, headers=headers, timeout=5)
            data = resp.json()
            if data.get("photos"):
                img_url = data["photos"][0]["src"]["large"]
                img_resp = requests.get(img_url, timeout=10)
                return BytesIO(img_resp.content)
        except Exception: pass
    # 2. Try Unsplash
    if UNSPLASH_ACCESS_KEY:
        try:
            url = f"https://api.unsplash.com/search/photos?query={query}&per_page=1&client_id={UNSPLASH_ACCESS_KEY}"
            resp = requests.get(url, timeout=5)
            data = resp.json()
            if data.get("results"):
                img_url = data["results"][0]["urls"]["regular"]
                img_resp = requests.get(img_url, timeout=10)
                return BytesIO(img_resp.content)
        except Exception: pass
    return None

def hex_to_rgb(hex_str: str) -> Tuple[int, int, int]:
    hex_str = hex_str.lstrip('#')
    return tuple(int(hex_str[i:i+2], 16) for i in (0, 2, 4))

# ==========================================
# 2. GAMMA-STYLE UI RENDERER
# ==========================================
class GammaRenderer:
    def __init__(self, design: Dict[str, Any]):
        self.prs = Presentation()
        self.width = 10
        self.height = 5.625
        self.prs.slide_width = Inches(self.width)
        self.prs.slide_height = Inches(self.height)
        
        # Design Tokens
        self.bg = hex_to_rgb(design.get("bg_color", "#05050A"))
        self.text = hex_to_rgb(design.get("text_color", "#FFFFFF"))
        self.accent = hex_to_rgb(design.get("accent_color", "#00F0FF"))
        self.muted = hex_to_rgb(design.get("muted_color", "#888899"))
        self.font_h = design.get("font_heading", "Inter Black")
        self.font_b = design.get("font_body", "Inter")

    def _set_bg(self, slide, with_accent=True):
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(*self.bg)
        
        if with_accent:
            # Sophisticated background geometry (UI Sidebar effect)
            bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(0.04), Inches(5.625))
            bar.fill.solid()
            bar.fill.fore_color.rgb = RGBColor(*self.accent)
            bar.line.visible = False

    def _add_text(self, shape, text, size=24, color=None, is_h=False, align=PP_ALIGN.LEFT):
        tf = shape.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = str(text)
        run = p.runs[0]
        run.font.size = Pt(size)
        run.font.name = self.font_h if is_h else self.font_b
        run.font.bold = is_h
        run.font.color.rgb = RGBColor(*(color or self.text))
        p.alignment = align
        return p

    def add_title_slide(self, title: str, subtitle: str):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self._set_bg(slide, with_accent=False)
        
        # Centered High-Impact Title
        box = slide.shapes.add_textbox(Inches(1), Inches(1.8), Inches(8), Inches(2))
        self._add_text(box, title.upper(), size=58, color=self.accent, is_h=True, align=PP_ALIGN.CENTER)
        
        sub_box = slide.shapes.add_textbox(Inches(1), Inches(3.2), Inches(8), Inches(1))
        self._add_text(sub_box, subtitle, size=22, color=self.muted, align=PP_ALIGN.CENTER)
        
        # Accent decoration
        line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(4), Inches(3.1), Inches(2), Inches(0.03))
        line.fill.solid()
        line.fill.fore_color.rgb = RGBColor(*self.accent)
        line.line.visible = False

    def add_content_slide(self, slide_data: Dict[str, Any]):
        layout = slide_data.get("layout", "BULLET").upper()
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self._set_bg(slide)
        
        # Heading (Modern Positioning)
        head_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(8.8), Inches(0.8))
        self._add_text(head_box, slide_data.get("heading", ""), size=34, color=self.accent, is_h=True)
        
        img_q = slide_data.get("image_query")
        img_stream = search_image(img_q) if img_q else None
        
        if layout == "GRID":
            self._render_grid(slide, slide_data.get("bullet_points", []))
        elif layout == "SPLIT" and img_stream:
            slide.shapes.add_picture(img_stream, Inches(0.6), Inches(1.3), height=Inches(3.8))
            txt_box = slide.shapes.add_textbox(Inches(5.2), Inches(1.3), Inches(4.2), Inches(3.8))
            self._render_bullets(txt_box, slide_data.get("bullet_points", []))
        elif layout == "HERO" and img_stream:
            # Full bleed with text card
            slide.shapes.add_picture(img_stream, Inches(0), Inches(0), width=Inches(10))
            overlay = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(5.5), Inches(0), Inches(4.5), Inches(5.625))
            overlay.fill.solid()
            overlay.fill.fore_color.rgb = RGBColor(*self.bg)
            overlay.fill.transparency = 0.2
            overlay.line.visible = False
            txt_box = slide.shapes.add_textbox(Inches(5.8), Inches(1), Inches(3.5), Inches(4))
            self._render_bullets(txt_box, slide_data.get("bullet_points", []), font_size=20)
        elif "table" in slide_data:
            self._render_table(slide, slide_data["table"])
        else:
            w = 5 if img_stream else 8.5
            txt_box = slide.shapes.add_textbox(Inches(0.6), Inches(1.3), Inches(w), Inches(3.8))
            self._render_bullets(txt_box, slide_data.get("bullet_points", []))
            if img_stream:
                pic = slide.shapes.add_picture(img_stream, Inches(5.8), Inches(1.3), width=Inches(3.6))

    def _render_bullets(self, box, bullets, font_size=18):
        tf = box.text_frame
        tf.word_wrap = True
        # Dynamic Scaling
        total = sum(len(str(b)) for b in bullets)
        if total > 500 or len(bullets) > 6: font_size = 14
        
        for i, b in enumerate(bullets):
            p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
            p.text = str(b).replace(">>", "›")
            p.font.size = Pt(font_size)
            p.font.name = self.font_b
            p.font.color.rgb = RGBColor(*self.text)
            p.space_before = Pt(8)

    def _render_grid(self, slide, items):
        count = min(3, len(items))
        if count == 0: return
        w = (9.0 / count) - 0.2
        for i in range(count):
            left = 0.6 + (i * (w + 0.2))
            # Card Shape
            rect = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(left), Inches(1.3), Inches(w), Inches(3.5))
            rect.fill.solid()
            rect.fill.fore_color.rgb = RGBColor(*self.bg)
            rect.line.color.rgb = RGBColor(*self.accent)
            rect.line.width = Pt(1)
            
            txt_box = slide.shapes.add_textbox(Inches(left + 0.1), Inches(1.4), Inches(w - 0.2), Inches(3.3))
            self._render_bullets(txt_box, [items[i]], font_size=14)

    def _render_table(self, slide, data):
        h, r = data.get("headers", []), data.get("rows", [])
        if not h or not r: return
        tbl = slide.shapes.add_table(len(r)+1, len(h), Inches(0.6), Inches(1.3), Inches(8.8), Inches(3.8)).table
        for c, hdr in enumerate(h):
            cell = tbl.cell(0, c)
            cell.text = str(hdr).upper()
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(*self.accent)
            p = cell.text_frame.paragraphs[0]
            p.font.bold = True
            p.font.color.rgb = RGBColor(*self.bg)
        for row_idx, row in enumerate(r):
            for col_idx, val in enumerate(row):
                if col_idx < len(h):
                    cell = tbl.cell(row_idx+1, col_idx)
                    cell.text = str(val)
                    cell.text_frame.paragraphs[0].font.color.rgb = RGBColor(*self.text)

    def save(self, path: str):
        self.prs.save(path)

# ==========================================
# 3. COORDINATOR
# ==========================================
def run_ppt_generation(prompt: str, context: str, user_memory: Dict[str, Any] = None, slide_count: int = 7) -> str:
    model = ModelFactory.create(model_platform=ModelPlatformType.OPENAI, model_type="gpt-4o", model_config_dict={"temperature": 0.8})
    
    # Extract memory snippets
    memory_str = ""
    if user_memory:
        patterns = ", ".join(user_memory.get("behavioralPatterns", []))
        mission = user_memory.get("longTermPlan", {}).get("mission", "")
        memory_str = f"\n[USER PREFERENCES]: {patterns}\n[USER MISSION]: {mission}"

    sys_prompt = f"""You are a Lead Designer for Gamma AI. Create a high-end presentation blueprint for: "{prompt}".
{memory_str}

**CONTEXTUAL HANDLING:**
- If the [CONTEXT] contains an existing presentation structure, you are performing a **REVISION/EVALUATION**.
- Respect existing slides but apply the requested changes or additions.
- Ensure the narrative flow remains consistent.

**DESIGN REQUIREMENTS:**
1. **Design System**: Suggest a professional HEX palette (BG, Text, Accent, Muted) based on topic mood.
2. **Layout Mix**: High variety. Use `GRID` for cards, `SPLIT` for images, `HERO` for impact, and `BULLET` for detail.
3. **Imagery**: Mandatory `image_query` for every slide. Be specific.

**OUTPUT SCHEMA (JSON ONLY):**
{{
  "design": {{ "bg_color": "#HEX", "text_color": "#HEX", "accent_color": "#HEX", "muted_color": "#HEX", "font_heading": "Arial Black", "font_body": "Arial" }},
  "slides": [
    {{ "layout": "TITLE", "title": "...", "subtitle": "..." }},
    {{ "layout": "GRID", "heading": "Key Pillars", "bullet_points": ["Pillar 1", "Pillar 2", "Pillar 3"], "image_query": "..." }},
    ... (total {slide_count} slides)
  ]
}}"""
    msg = BaseMessage.make_assistant_message(role_name="GammaDesigner", content=sys_prompt)
    resp = ChatAgent(msg, model=model).step(BaseMessage.make_user_message(role_name="User", content=f"Context: {context}\nTopic: {prompt}")).msg.content
    return resp.replace("```json", "").replace("```", "").strip()

async def generate_ppt(prompt: str, context: str, user_memory: Dict[str, Any] = None, slide_count: int = 7) -> Dict[str, Any]:
    raw = run_ppt_generation(prompt, context, user_memory, slide_count)
    try:
        data = json.loads(raw)
        design = data.get("design", {})
        slides = data.get("slides", [])
    except:
        return {"success": False, "error": "Invalid blueprint"}

    renderer = GammaRenderer(design)
    for i, slide in enumerate(slides):
        if i == 0 or slide.get("layout") == "TITLE":
            renderer.add_title_slide(slide.get("title", prompt), slide.get("subtitle", ""))
        else:
            renderer.add_content_slide(slide)
            
    fd, path = tempfile.mkstemp(suffix=".pptx")
    os.close(fd)
    renderer.save(path)
    with open(path, "rb") as f:
        buf = f.read()
    b64 = base64.b64encode(buf).decode('utf-8')
    os.remove(path)
    return {"slides": slides, "base64_file": b64}
